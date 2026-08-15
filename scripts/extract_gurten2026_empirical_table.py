#!/usr/bin/env python3
"""Extract Gurten et al. (2026) Supplement 1 empirical PCR results.

The source DOCX is retained unchanged under data/external/gurten2026. This
script converts its single five-column table to analysis-ready CSV files while
keeping the authors' taxon strings and yes/no outcomes verbatim.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    ROOT
    / "data"
    / "external"
    / "gurten2026"
    / "supplement_1_empirical_amplification.docx"
)
OUTPUT = ROOT / "data" / "derived" / "beeprime_empirical_validation.csv"
GENUS_OUTPUT = ROOT / "data" / "derived" / "beeprime_empirical_genus_summary.csv"
MANIFEST_OUTPUT = ROOT / "data" / "derived" / "beeprime_empirical_manifest.csv"

SOURCE_DOI = "10.3897/mbmg.10.183708"
SUPPLEMENT_DOI = "10.3897/mbmg.10.183708.suppl1"


def clean(value: str) -> str:
    return " ".join(value.replace("\n", " ").split())


def inferred_rank(taxon: str) -> str:
    tokens = taxon.split()
    if taxon.endswith(("idae", "inae")):
        return "family_or_subfamily"
    if len(tokens) >= 2 and tokens[1].lower().rstrip(".") == "sp":
        return "genus"
    if len(tokens) >= 2 and tokens[0][0:1].isupper():
        return "species_or_complex"
    return "higher_taxon"


def inferred_genus(taxon: str) -> str:
    match = re.match(r"^([A-Z][A-Za-z-]+)(?:\s|$)", taxon)
    if not match or taxon.endswith(("idae", "inae")):
        return ""
    return match.group(1)


def main() -> None:
    document = Document(SOURCE)
    if len(document.tables) != 1:
        raise RuntimeError(f"Expected one table, found {len(document.tables)}")

    table = document.tables[0]
    rows: list[dict[str, str]] = []
    current_order = ""

    for row_number, row in enumerate(table.rows[2:], start=3):
        values = [clean(cell.text) for cell in row.cells]
        order, common_name, taxon, target, detected = values

        # Rows containing only an order name are section headers.
        if order and not any((common_name, taxon, target, detected)):
            current_order = order
            continue
        if not taxon:
            continue

        target = target.lower()
        detected = detected.lower()
        if target not in {"yes", "no"} or detected not in {"yes", "no"}:
            raise RuntimeError(
                f"Unexpected yes/no values on source table row {row_number}: "
                f"target={target!r}, detected={detected!r}"
            )

        rows.append(
            {
                "source_table_row": str(row_number),
                "order": current_order,
                "common_name": common_name,
                "taxon": taxon,
                "inferred_rank": inferred_rank(taxon),
                "genus": inferred_genus(taxon),
                "target": target,
                "detected": detected,
                "target_bool": str(target == "yes").upper(),
                "detected_bool": str(detected == "yes").upper(),
                "source_doi": SOURCE_DOI,
                "source_supplement_doi": SUPPLEMENT_DOI,
            }
        )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)

    target_rows = [row for row in rows if row["target"] == "yes"]
    # The machine-readable Supplement 1 table contains 37 target rows and 32
    # detections, whereas the article text reports 38 and 33. Preserve both
    # counts explicitly instead of silently forcing them to agree.
    if len(target_rows) != 37:
        raise RuntimeError(f"Expected 37 target rows in Supplement 1, found {len(target_rows)}")
    detected_targets = sum(row["detected"] == "yes" for row in target_rows)
    if detected_targets != 32:
        raise RuntimeError(
            f"Expected 32 detected target rows in Supplement 1, found {detected_targets}"
        )

    genus_counts: dict[str, dict[str, int]] = {}
    for row in target_rows:
        genus = row["genus"] or "Unresolved"
        stats = genus_counts.setdefault(
            genus, {"n_empirical_tested": 0, "n_empirical_detected": 0}
        )
        stats["n_empirical_tested"] += 1
        stats["n_empirical_detected"] += int(row["detected"] == "yes")

    genus_rows = []
    for genus, stats in sorted(genus_counts.items()):
        tested = stats["n_empirical_tested"]
        detected = stats["n_empirical_detected"]
        genus_rows.append(
            {
                "genus": genus,
                "n_empirical_tested": tested,
                "n_empirical_detected": detected,
                "empirical_detection_fraction": detected / tested,
                "source_doi": SOURCE_DOI,
            }
        )

    with GENUS_OUTPUT.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(genus_rows[0]))
        writer.writeheader()
        writer.writerows(genus_rows)

    manifest_rows = [
        ("source_article", "Gurten et al. 2026"),
        ("source_doi", SOURCE_DOI),
        ("source_supplement_doi", SUPPLEMENT_DOI),
        ("article_text_target_bee_extracts", "38"),
        ("article_text_target_bee_detected", "33"),
        ("article_text_detection_fraction", f"{33 / 38:.8f}"),
        ("supplement_table_target_rows", str(len(target_rows))),
        ("supplement_table_target_detected", str(detected_targets)),
        (
            "supplement_table_detection_fraction",
            f"{detected_targets / len(target_rows):.8f}",
        ),
        (
            "count_note",
            "Article text reports 38/33; Supplement 1 lists 37/32. "
            "Both imply five listed target failures. The app shows both sources.",
        ),
    ]
    with MANIFEST_OUTPUT.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(("field", "value"))
        writer.writerows(manifest_rows)

    print(
        f"Wrote {len(rows)} empirical rows. Supplement 1 lists target bees "
        f"{detected_targets}/{len(target_rows)}; article text reports 33/38."
    )


if __name__ == "__main__":
    main()
